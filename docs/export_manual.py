"""
@author zhangyueting
@date 2026-06-12
将 用户手册.md 转换为 Word(.docx) 与自包含 HTML（供导出 PDF）。
仅处理本手册使用到的 Markdown 结构：H1-H3、段落、有序/无序列表、
表格、图片、引用块、代码块、分割线、行内 **粗体** 与 `行内代码`。
"""
import os
import re
import base64
import html as htmllib
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD = os.path.join(ROOT, "用户手册.md")
DOCX = os.path.join(ROOT, "用户手册.docx")
HTML = os.path.join(ROOT, "用户手册.html")

with open(MD, "r", encoding="utf-8") as f:
    lines = f.read().split("\n")


# ---------------- 公共：行内解析 ----------------
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+?`)")


def inline_segments(text):
    """返回 [(type, content)]，type in {'b','code','t'}"""
    segs = []
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            segs.append(("b", part[2:-2]))
        elif part.startswith("`") and part.endswith("`"):
            segs.append(("code", part[1:-1]))
        else:
            segs.append(("t", part))
    return segs


# ===================================================================
#  Word 生成
# ===================================================================
def build_docx():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)
    # 中文字体
    from docx.oxml.ns import qn
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    def add_inline(p, text):
        for typ, content in inline_segments(text):
            run = p.add_run(content)
            run.font.name = "Microsoft YaHei"
            run.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            if typ == "b":
                run.bold = True
            elif typ == "code":
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            p.paragraph_format.left_indent = Inches(0.2)
            continue

        # 图片
        m = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
        if m:
            alt, src = m.group(1), m.group(2)
            img_path = os.path.join(ROOT, src.replace("/", os.sep))
            if os.path.exists(img_path):
                with Image.open(img_path) as im:
                    w_px = im.width
                # 限制最大宽度 6.2 英寸
                width = min(Inches(6.2), Inches(w_px / 96))
                doc.add_picture(img_path, width=width)
                last = doc.paragraphs[-1]
                last.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if alt:
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = cap.add_run("图：" + alt)
                    r.italic = True
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(0x90, 0x90, 0x90)
            i += 1
            continue

        # 标题
        if stripped.startswith("#"):
            lvl = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[lvl:].strip()
            doc.add_heading(text, level=min(lvl, 4))
            i += 1
            continue

        # 分割线
        if stripped in ("---", "***", "___"):
            i += 1
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < n and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for j, h in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.paragraphs[0].text = ""
                add_inline(cell.paragraphs[0], h)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                cells = table.add_row().cells
                for j, val in enumerate(row):
                    if j < len(cells):
                        cells[j].paragraphs[0].text = ""
                        add_inline(cells[j].paragraphs[0], val)
            doc.add_paragraph()
            continue

        # 引用块
        if stripped.startswith(">"):
            text = stripped.lstrip(">").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_inline(p, text)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
            i += 1
            continue

        # 无序列表
        if re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^[-*]\s+", "", stripped))
            i += 1
            continue

        # 有序列表
        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    doc.save(DOCX)
    print("DOCX_OK", DOCX)


# ===================================================================
#  HTML 生成（图片转 base64，便于打印为 PDF）
# ===================================================================
def build_html():
    import markdown

    raw = "\n".join(lines)

    def repl_img(m):
        alt, src = m.group(1), m.group(2)
        p = os.path.join(ROOT, src.replace("/", os.sep))
        if os.path.exists(p):
            ext = os.path.splitext(p)[1].lstrip(".").lower()
            with open(p, "rb") as fh:
                b64 = base64.b64encode(fh.read()).decode()
            return f'![{alt}](data:image/{ext};base64,{b64})'
        return m.group(0)

    raw = re.sub(r"!\[(.*?)\]\((.*?)\)", repl_img, raw)
    body = markdown.markdown(raw, extensions=["tables", "fenced_code", "toc"])

    css = """
    body{font-family:'Microsoft YaHei','PingFang SC',sans-serif;max-width:900px;margin:0 auto;padding:32px;color:#222;line-height:1.7;}
    h1{font-size:26px;border-bottom:3px solid #4c8bf5;padding-bottom:8px;}
    h2{font-size:20px;border-bottom:1px solid #ddd;padding-bottom:6px;margin-top:32px;}
    h3{font-size:16px;margin-top:24px;}
    table{border-collapse:collapse;width:100%;margin:12px 0;}
    th,td{border:1px solid #ccc;padding:6px 10px;font-size:14px;text-align:left;}
    th{background:#f0f4ff;}
    img{max-width:100%;border:1px solid #e0e0e0;border-radius:4px;box-shadow:0 1px 4px rgba(0,0,0,.08);margin:8px 0;}
    code{background:#f5f5f5;color:#c7254e;padding:2px 5px;border-radius:3px;font-family:Consolas,monospace;font-size:90%;}
    pre{background:#f6f8fa;padding:12px;border-radius:6px;overflow:auto;}
    pre code{background:none;color:#333;}
    blockquote{border-left:4px solid #4c8bf5;background:#f8faff;margin:8px 0;padding:8px 14px;color:#555;}
    a{color:#4c8bf5;}
    """
    full = f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">
<title>TCMP 用户手册</title><style>{css}</style></head><body>{body}</body></html>"""
    with open(HTML, "w", encoding="utf-8") as fh:
        fh.write(full)
    print("HTML_OK", HTML)


if __name__ == "__main__":
    build_docx()
    build_html()
